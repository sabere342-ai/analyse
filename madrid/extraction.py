"""Bounded local extraction. Unknown POS layouts are never claimed as verified adapters."""

from __future__ import annotations

import csv
import io
import json
import os
import resource
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

MAX_ROWS = 10000
MAX_BYTES = min(20 * 1024 * 1024, int(os.environ.get("MADRID_MAX_UPLOAD_BYTES", 20 * 1024 * 1024)))
if MAX_BYTES <= 0:
    raise ValueError("invalid_upload_limit")


def validate_file(path: Path, suffix: str) -> None:
    data = path.read_bytes()
    if not data or len(data) > MAX_BYTES:
        raise ValueError("empty_or_oversized_file")
    if suffix not in {".csv", ".pdf", ".xls", ".xlsx"}:
        raise ValueError("unsupported_file_type")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("invalid_pdf_signature")
    if suffix == ".xls" and not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise ValueError("invalid_xls_signature")
    if suffix == ".xlsx":
        if not zipfile.is_zipfile(path):
            raise ValueError("invalid_xlsx_signature")
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
                raise ValueError("workbook_expansion_limit")
            if any(i.file_size > 1024 * max(i.compress_size, 1) or i.flag_bits & 1 for i in infos):
                raise ValueError("unsafe_workbook_archive")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("not_an_xlsx_workbook")
            if any("vbaProject" in n or "externalLinks/" in n for n in names):
                raise ValueError("macros_or_external_links_prohibited")
    if suffix == ".csv" and (b"\x00" in data or data.startswith((b"%PDF", b"PK\x03\x04"))):
        raise ValueError("not_a_text_csv")


def extract(path: Path, suffix: str, options: dict[str, Any]) -> list[dict[str, Any]]:
    validate_file(path, suffix)
    result: list[dict[str, Any]] = []

    def add(
        cells: list[Any], location: dict[str, Any], numeric: list[int] | None = None, formula: bool = False
    ) -> None:
        if not any(c is not None and str(c).strip() for c in cells):
            return
        if len(cells) > 100 or len(result) >= MAX_ROWS:
            raise ValueError("row_or_column_limit")
        if any(len(str(c or "")) > 2000 for c in cells):
            raise ValueError("cell_text_limit")
        result.append(
            {
                "cells": [
                    ""
                    if c is None
                    else c.date().isoformat()
                    if isinstance(c, datetime)
                    else c.isoformat()
                    if isinstance(c, date)
                    else str(c)
                    for c in cells
                ],
                "location": location,
                "numeric_columns": numeric or [],
                "formula": formula,
            }
        )

    if suffix == ".csv":
        text = path.read_bytes().decode(options.get("encoding", "utf-8-sig"), errors="strict")
        delimiter = options.get("delimiter", "auto")
        if delimiter == "auto":
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
            except csv.Error as exc:
                raise ValueError("csv_delimiter_requires_selection") from exc
        for i, row in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), 1):
            add(row, {"row": i})
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if len(book.worksheets) > 20:
            raise ValueError("worksheet_limit")
        if set(options.get("sheets", [])) - set(book.sheetnames):
            book.close()
            raise ValueError("selected_worksheet_not_found")
        for sheet in book:
            if options.get("sheets") and sheet.title not in options["sheets"]:
                continue
            for i, row in enumerate(sheet.iter_rows(), 1):
                add(
                    [c.value for c in row],
                    {"sheet": sheet.title, "row": i},
                    [j for j, c in enumerate(row) if c.data_type == "n" and c.value is not None],
                    any(c.data_type == "f" for c in row),
                )
        book.close()
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        if book.nsheets > 20:
            raise ValueError("worksheet_limit")
        if set(options.get("sheets", [])) - set(book.sheet_names()):
            book.release_resources()
            raise ValueError("selected_worksheet_not_found")
        for sheet in book.sheets():
            if options.get("sheets") and sheet.name not in options["sheets"]:
                continue
            for i in range(sheet.nrows):
                cells = sheet.row(i)
                add(
                    [
                        int(c.value) if c.ctype == 2 and float(c.value).is_integer() else c.value
                        for c in cells
                    ],
                    {"sheet": sheet.name, "row": i + 1},
                    [j for j, c in enumerate(cells) if c.ctype == 2],
                )
        book.release_resources()
    else:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > 200:
                raise ValueError("pdf_page_limit")
            from madrid.pos_pdf import extract_pos

            profiled = extract_pos(pdf, path, options)
            if profiled is not None:
                return profiled
            for p, page in enumerate(pdf.pages, 1):
                tables = page.extract_tables()
                if not tables:
                    raise ValueError("pdf_layout_requires_adapter_no_table_or_text_layer")
                for table_index, table in enumerate(tables, 1):
                    for i, pdf_row in enumerate(table, 1):
                        add(pdf_row, {"page": p, "table": table_index, "row": i})
    if not result:
        raise ValueError("no_rows_in_selected_source")
    return result


if __name__ == "__main__":
    resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
    resource.setrlimit(resource.RLIMIT_CPU, (30, 30))
    resource.setrlimit(resource.RLIMIT_FSIZE, (128 * 1024**2, 128 * 1024**2))
    try:
        payload = json.load(sys.stdin)
        output = extract(Path(payload["path"]), payload["suffix"], payload["options"])
        print(json.dumps({"rows": output}, ensure_ascii=False))
    except Exception as exc:
        # Error types are safe; never echo arbitrary document text or library stack traces.
        allowed = str(exc) if isinstance(exc, ValueError) and str(exc).isascii() else type(exc).__name__
        print(json.dumps({"error": allowed[:120]}))
        sys.exit(1)
