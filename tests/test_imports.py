import csv
import io
import json
from concurrent.futures import ThreadPoolExecutor

import pytest
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from madrid.contracts import code_value, date_value, decimal_value
from madrid.extraction import MAX_BYTES
from madrid.ingestion import transition


def csv_bytes(rows, encoding="utf-8-sig", delimiter=","):
    stream = io.StringIO(newline="")
    csv.writer(stream, delimiter=delimiter).writerows(rows)
    return stream.getvalue().encode(encoding)


def upload(client, rows, kind="snapshot", start="2026-10-24", end=None, expected=None, **options):
    spec = {
        "report_type": kind,
        "start": start,
        "end": end or start,
        "complete": True,
        "expected": expected or {},
        **options,
    }
    response = client.post(
        "/api/imports",
        files={"file": ("synthetic.csv", csv_bytes(rows), "text/csv")},
        data={"spec": json.dumps(spec)},
    )
    assert response.status_code == 200, response.text
    return response.json()


def accept(client, batch):
    result = client.post(f"/api/batches/{batch['id']}/action", json={"action": "accept"})
    assert result.status_code == 200, result.text
    return result.json()


SNAPSHOT = [
    ["barcode", "name", "quantity", "reserved", "available"],
    ["000366", "تي شيرت Basic", "10", "2", "8"],
    ["TOTAL", "", "10", "2", "8"],
]


def test_roundtrip_duplicate_overlap_and_rollback(client):
    batch = upload(client, SNAPSHOT, reservation_verified=True)
    assert batch["issues"] == []
    assert batch["rows"][1]["data"]["barcode"] == "000366"
    assert batch["reconciliation"]["checks"]["quantity"]["matches"]
    before = client.get("/api/analytics?start=2026-10-01&end=2026-10-24").json()
    assert before["metrics"]["stock"]["value"] is None
    accept(client, batch)
    duplicate = upload(client, SNAPSHOT, reservation_verified=True)
    assert duplicate["id"] == batch["id"] and duplicate["status"] == "accepted"
    other = upload(client, SNAPSHOT, reservation_verified=True, note="new interpretation")
    result = client.post(f"/api/batches/{other['id']}/action", json={"action": "accept"})
    assert result.status_code == 409
    after = client.get("/api/analytics?start=2026-10-01&end=2026-10-24").json()
    assert after["metrics"]["stock"]["value"] == "10.00"
    assert after["products"][0]["available"] == "8"
    result = client.post(
        f"/api/batches/{batch['id']}/action", json={"action": "rollback", "reason": "synthetic correction"}
    )
    assert result.status_code == 200
    assert (
        client.get("/api/analytics?start=2026-10-01&end=2026-10-24").json()["metrics"]["stock"]["value"]
        is None
    )
    accept(client, other)
    assert any(a["action"] == "rollback" for a in client.get("/api/audit").json())


def test_import_comparison_and_new_snapshot_update_without_adding_balances(client):
    first = upload(client, SNAPSHOT, reservation_verified=True)
    preview = client.get(f"/api/batches/{first['id']}/comparison").json()
    assert preview["new_products"] == 1 and preview["can_accept"]
    accept(client, first)
    changed = upload(
        client,
        [
            ["barcode", "name", "quantity"],
            ["000366", "Existing item", "7"],
            ["000999", "New item", "2"],
            ["TOTAL", "", "9"],
        ],
        start="2026-10-25",
    )
    comparison = client.get(f"/api/batches/{changed['id']}/comparison").json()
    assert comparison["existing_products"] == comparison["new_products"] == 1
    assert comparison["changed_stock_products"] == 1
    assert comparison["stock_changes"][0]["difference"] == "-3"
    assert comparison["previous_snapshot_date"] == "2026-10-24"
    assert comparison["can_accept"] and not comparison["overlaps"]
    query = "/api/analytics?start=2026-10-01&end=2026-10-25&stock_asof=2026-10-25"
    assert client.get(query).json()["metrics"]["stock"]["value"] == "10.00"
    accept(client, changed)
    assert client.get(query).json()["metrics"]["stock"]["value"] == "9.00"
    correction = upload(client, SNAPSHOT, note="review correction")
    blocked = client.get(f"/api/batches/{correction['id']}/comparison").json()
    assert not blocked["can_accept"]
    assert blocked["overlaps"][0]["id"] == first["id"]
    assert (
        client.post(f"/api/batches/{correction['id']}/action", json={"action": "accept"}).status_code == 409
    )


def test_import_comparison_keeps_other_coverage_scopes_separate(client):
    accept(client, upload(client, SNAPSHOT, scope="other-store", reservation_verified=True))
    incoming = upload(client, SNAPSHOT, scope="store", reservation_verified=True)
    comparison = client.get(f"/api/batches/{incoming['id']}/comparison").json()
    assert comparison["new_products"] == 1 and comparison["existing_products"] == 0
    assert not comparison["overlaps"] and comparison["can_accept"]


@pytest.mark.parametrize(
    "rows,expected_issue",
    [
        (
            [["barcode", "name", "quantity"], ["", "منتج", "10"], ["TOTAL", "", "10"]],
            "row_issues_require_corrected_reimport",
        ),
        (
            [["barcode", "name", "quantity"], ["0001", "منتج", "9"], ["TOTAL", "", "10"]],
            "source_total_mismatch",
        ),
        (
            [["barcode", "name", "quantity"], ["0001", "منتج", "9"]],
            "independent_source_total_required:quantity",
        ),
        (
            [
                ["barcode", "name", "quantity"],
                ["0001", "منتج", "5"],
                ["0001", "منتج", "5"],
                ["TOTAL", "", "10"],
            ],
            "row_issues_require_corrected_reimport",
        ),
    ],
)
def test_invalid_rows_never_accepted(client, rows, expected_issue):
    batch = upload(client, rows)
    assert expected_issue in batch["issues"]
    assert client.post(f"/api/batches/{batch['id']}/action", json={"action": "accept"}).status_code == 409


def test_explicit_expected_cannot_hide_wrong_footer(client):
    rows = [["barcode", "name", "quantity"], ["0001", "منتج", "9"], ["TOTAL", "", "10"]]
    batch = upload(client, rows, expected={"quantity": "9"})
    assert "source_total_mismatch" in batch["issues"]


def test_headers_subtotals_continuation_and_mixed_text(client):
    rows = [
        ["barcode", "name", "quantity"],
        ["", "تي شيرت Basic", "5"],
        ["000366", "", ""],
        ["SUBTOTAL", "", "5"],
        ["barcode", "name", "quantity"],
        ["000369", "قميص Oxford", "3"],
        ["TOTAL", "", "8"],
    ]
    batch = upload(client, rows)
    assert batch["issues"] == []
    assert batch["reconciliation"]["details"] == 2
    assert batch["rows"][1]["continuation"] == {"row": 3}
    assert batch["rows"][1]["data"]["barcode"] == "000366"


def test_legitimate_repeated_sales_survive(client):
    rows = [
        ["barcode", "name", "sold", "returned"],
        ["000366", "A", "2", "0"],
        ["000366", "A", "2", "0"],
        ["TOTAL", "", "4", "0"],
    ]
    batch = upload(client, rows, "sales", "2026-10-01", "2026-10-23")
    assert batch["issues"] == [] and batch["reconciliation"]["details"] == 2
    accept(client, batch)
    assert (
        client.get("/api/analytics?start=2026-10-01&end=2026-10-23").json()["metrics"]["sold"]["value"]
        == "4.00"
    )


@pytest.mark.parametrize("encoding,delimiter", [("utf-8-sig", ";"), ("cp1256", "\t")])
def test_csv_encodings_and_delimiters(client, encoding, delimiter):
    body = csv_bytes(SNAPSHOT, encoding, delimiter)
    spec = {
        "report_type": "snapshot",
        "start": "2026-10-24",
        "end": "2026-10-24",
        "complete": True,
        "encoding": encoding,
        "delimiter": delimiter,
    }
    result = client.post(
        "/api/imports", files={"file": ("arabic.csv", body)}, data={"spec": json.dumps(spec)}
    )
    assert result.status_code == 200, result.text
    assert result.json()["issues"] == []


def test_xlsx_multiple_sheets_numeric_codes_and_formulas(client):
    book = Workbook()
    sheet = book.active
    sheet.append(["barcode", "name", "quantity"])
    sheet.append(["000366", "قميص", "3"])
    other = book.create_sheet("Second")
    other.append(["barcode", "name", "quantity"])
    other.append([369, "قميص", "2"])
    other.append(["000377", "حذاء", "=1+1"])
    stream = io.BytesIO()
    book.save(stream)
    spec = {
        "report_type": "snapshot",
        "start": "2026-10-24",
        "end": "2026-10-24",
        "complete": True,
        "expected": {"quantity": "7"},
    }
    result = client.post(
        "/api/imports", files={"file": ("synthetic.xlsx", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    batch = result.json()
    assert result.status_code == 200, result.text
    issues = [issue for row in batch["rows"] for issue in row["issues"]]
    assert "numeric_barcode_requires_text_source" in issues
    assert "workbook_formula_requires_review" in issues
    assert batch["rows"][1]["data"]["barcode"] == "000366"
    assert batch["rows"][3]["raw"]["location"]["sheet"] == "Second"


def test_legacy_xls(client):
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("Synthetic")
    for y, row in enumerate(SNAPSHOT):
        for x, value in enumerate(row):
            sheet.write(y, x, value)
    stream = io.BytesIO()
    book.save(stream)
    spec = {"report_type": "snapshot", "start": "2026-10-24", "end": "2026-10-24", "complete": True}
    response = client.post(
        "/api/imports", files={"file": ("synthetic.xls", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert response.status_code == 200, response.text
    assert response.json()["issues"] == []


def test_pdf_multipage_tables(client):
    buffer = io.BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=(450, 240), topMargin=15, bottomMargin=15)
    table = Table(
        [
            ["barcode", "name", "quantity"],
            *[[f"{i:06}", "Synthetic item", "1"] for i in range(1, 26)],
            ["TOTAL", "", "25"],
        ],
        repeatRows=1,
    )
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    document.build([table])
    spec = {"report_type": "snapshot", "start": "2026-10-24", "end": "2026-10-24", "complete": True}
    response = client.post(
        "/api/imports", files={"file": ("synthetic.pdf", buffer.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert response.status_code == 200, response.text
    batch = response.json()
    assert batch["issues"] == []
    assert batch["reconciliation"]["details"] == 25
    assert max(r["raw"]["location"]["page"] for r in batch["rows"]) > 1


@pytest.mark.parametrize(
    "filename,content",
    [
        ("broken.pdf", b"%PDF-broken"),
        ("fake.xlsx", b"not zip"),
        ("file.exe", b"anything"),
        ("fake.csv", b"%PDF-binary"),
    ],
)
def test_bad_files_safe_failure(client, filename, content):
    spec = {"report_type": "snapshot", "start": "2026-10-24", "end": "2026-10-24"}
    response = client.post(
        "/api/imports", files={"file": (filename, content)}, data={"spec": json.dumps(spec)}
    )
    assert response.status_code == 422
    assert "Traceback" not in response.text
    assert client.get("/api/batches").json() == []


def test_oversized_upload_rejected(client):
    response = client.post(
        "/api/imports", files={"file": ("huge.csv", b"x" * (MAX_BYTES + 100000))}, data={"spec": "{}"}
    )
    assert response.status_code == 413


def test_parser_version_changes_interpretation_not_accepted_coverage(client, monkeypatch):
    from madrid import ingestion

    a = upload(client, SNAPSHOT)
    accept(client, a)
    monkeypatch.setattr(ingestion, "PARSER_VERSION", "synthetic-next-version")
    b = upload(client, SNAPSHOT)
    assert a["id"] != b["id"]
    assert client.post(f"/api/batches/{b['id']}/action", json={"action": "accept"}).status_code == 409


def test_concurrent_accept_only_one_coverage(client):
    a = upload(client, SNAPSHOT, note="a")
    b = upload(client, SNAPSHOT, note="b")

    def run(batch_id):
        with client.app.state.factory() as session:
            try:
                transition(session, batch_id, "accept", "")
                return "accepted"
            except ValueError:
                return "conflict"

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(run, [a["id"], b["id"]])) == ["accepted", "conflict"]


@pytest.mark.parametrize("source,expected", [("٠٠٠٣٦٦", "000366"), (" 000366 ", "000366")])
def test_code_normalization(source, expected):
    assert code_value(source) == expected


@pytest.mark.parametrize("source,expected", [("١٬٢٣٤٫٥٠", "1234.50"), ("(20.25)", "-20.25"), ("−3", "-3")])
def test_decimal_parsing(source, expected):
    assert str(decimal_value(source)) == expected


@pytest.mark.parametrize("source", ["1,25", "NaN", "1e20", "Infinity", "1.0000001"])
def test_ambiguous_numbers_rejected(source):
    with pytest.raises(ValueError):
        decimal_value(source)


def test_day_month_year_dates():
    assert date_value("٢٤/١٠/٢٠٢٦") == "2026-10-24"
    with pytest.raises(ValueError):
        date_value("31/02/2026")


def test_encrypted_pdf_and_macros_are_rejected(client):
    import zipfile

    from reportlab.lib.pdfencrypt import StandardEncryption
    from reportlab.pdfgen import canvas

    stream = io.BytesIO()
    pdf = canvas.Canvas(stream, encrypt=StandardEncryption("synthetic-encryption-password"))
    pdf.drawString(40, 40, "Synthetic protected report")
    pdf.save()
    spec = {"report_type": "snapshot", "start": "2026-10-24", "end": "2026-10-24"}
    response = client.post(
        "/api/imports", files={"file": ("encrypted.pdf", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert response.status_code == 422
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "dummy")
        archive.writestr("xl/workbook.xml", "dummy")
        archive.writestr("xl/vbaProject.bin", "not executable")
    response = client.post(
        "/api/imports", files={"file": ("unsafe.xlsx", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert response.status_code == 422 and response.json()["detail"] == "macros_or_external_links_prohibited"


def test_chunked_upload_limit_does_not_rely_on_content_length(client):
    request = client.build_request(
        "POST",
        "/api/imports",
        files={"file": ("large.csv", b"x" * (MAX_BYTES + 100000))},
        data={"spec": "{}"},
    )
    del request.headers["content-length"]
    request.headers["transfer-encoding"] = "chunked"
    response = client.send(request)
    assert response.status_code == 413


def test_user_mapping_and_sheet_selection(client):
    book = Workbook()
    sheet = book.active
    sheet.title = "Approved"
    sheet.append(["My code", "My name", "My stock"])
    sheet.append(["000366", "تجريبي", "10"])
    sheet.append(["TOTAL", "", "10"])
    other = book.create_sheet("Unselected")
    other.append(["unrelated source"])
    stream = io.BytesIO()
    book.save(stream)
    spec = {
        "report_type": "snapshot",
        "start": "2026-10-24",
        "end": "2026-10-24",
        "complete": True,
        "sheets": ["Approved"],
        "column_map": {"barcode": "My code", "name": "My name", "quantity": "My stock"},
    }
    response = client.post(
        "/api/imports", files={"file": ("mapped.xlsx", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert response.status_code == 200 and not response.json()["issues"]
    assert all(r["raw"]["location"]["sheet"] == "Approved" for r in response.json()["rows"])
    spec["sheets"] = ["Approved", "Misspelled"]
    response = client.post(
        "/api/imports", files={"file": ("mapped.xlsx", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert response.status_code == 422 and response.json()["detail"] == "selected_worksheet_not_found"
    assert len(client.get("/api/batches").json()) == 1


@pytest.mark.parametrize(
    "first_quantity,second_quantity,expected_issue",
    [("2", "3", None), ("2", "4", "source_total_mismatch"), ("1", "4", "source_total_mismatch")],
)
def test_workbook_reconciles_each_sheet_before_combining(
    client, first_quantity, second_quantity, expected_issue
):
    book = Workbook()
    first = book.active
    first.title = "First"
    for values in [
        ["barcode", "name", "quantity"],
        ["000001", "First item", first_quantity],
        ["TOTAL", "", "2"],
    ]:
        first.append(values)
    second = book.create_sheet("Second")
    for values in [
        ["quantity", "name", "barcode"],
        [second_quantity, "Second item", "000002"],
        ["3", "TOTAL", ""],
    ]:
        second.append(values)
    stream = io.BytesIO()
    book.save(stream)
    spec = {"report_type": "snapshot", "start": "2026-10-24", "end": "2026-10-24", "complete": True}
    result = client.post(
        "/api/imports", files={"file": ("sheets.xlsx", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert result.status_code == 200, result.text
    batch = result.json()
    assert batch["reconciliation"]["details"] == 2
    assert batch["reconciliation"]["checks"]["quantity"]["expected"] == "5"
    if expected_issue:
        assert expected_issue in batch["issues"]
        assert client.post(f"/api/batches/{batch['id']}/action", json={"action": "accept"}).status_code == 409
    else:
        assert batch["issues"] == []
        accept(client, batch)
        report = client.get("/api/analytics?start=2026-10-24&end=2026-10-24").json()
        assert report["metrics"]["stock"]["value"] == "5.00"


def test_worksheet_cannot_inherit_another_worksheets_header(client):
    book = Workbook()
    first = book.active
    first.title = "Stock"
    for values in [["barcode", "name", "quantity"], ["000001", "First item", "2"]]:
        first.append(values)
    book.create_sheet("No header").append(["000002", "Unmapped item", "3"])
    stream = io.BytesIO()
    book.save(stream)
    spec = {
        "report_type": "snapshot",
        "start": "2026-10-24",
        "end": "2026-10-24",
        "complete": True,
        "expected": {"quantity": "5"},
    }
    result = client.post(
        "/api/imports", files={"file": ("unmapped.xlsx", stream.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert result.status_code == 200, result.text
    batch = result.json()
    assert batch["rows"][-1]["kind"] == "unknown"
    assert batch["rows"][-1]["issues"] == ["header_mapping_required"]
    assert "row_issues_require_corrected_reimport" in batch["issues"]
    assert client.post(f"/api/batches/{batch['id']}/action", json={"action": "accept"}).status_code == 409


def test_concurrent_identical_uploads_share_atomic_source_and_batch(client):
    from madrid.contracts import ImportSpec
    from madrid.ingestion import stage

    spec = ImportSpec(report_type="snapshot", start="2026-10-24", end="2026-10-24", complete=True)

    def run(_):
        return stage(
            client.app.state.factory, client.app.state.directory, "synthetic.csv", csv_bytes(SNAPSHOT), spec
        ).id

    with ThreadPoolExecutor(max_workers=2) as pool:
        ids = list(pool.map(run, range(2)))
    assert len(set(ids)) == 1
    assert len(client.get("/api/batches").json()) == 1
    assert not list((client.app.state.directory / "originals").glob("*.tmp"))
