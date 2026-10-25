"""Real PostgreSQL integration; isolated schemas, synthetic rows, verified TLS."""

import hashlib
import io
import json
import os
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import OperationalError
from test_imports import SNAPSHOT, accept, csv_bytes, upload

from madrid.app import create_app
from madrid.backup import create_cloud_backup, restore_backup
from madrid.db import Audit, AuthThrottle, Batch, Original, make_engine, session_factory

PASSWORD = "synthetic-cloud-password-only"
INVITATION = "synthetic-cloud-invitation-only"
HEADERS = {"X-Requested-With": "Madrid"}


@pytest.fixture
def cloud(tmp_path, monkeypatch):
    url = os.environ.get("MADRID_TEST_DATABASE_URL")
    if not url:
        pytest.skip("Real PostgreSQL test database not configured; cloud persistence is not verified")
    engine = make_engine(tmp_path / "unused", url)
    schema = "mri_test_" + uuid.uuid4().hex
    with engine.begin() as connection:
        connection.execute(text(f'CREATE SCHEMA "{schema}"'))
    parsed = make_url(url)
    query = dict(parsed.query)
    query["options"] = "-csearch_path=" + schema
    monkeypatch.setenv("MADRID_DATABASE_URL", parsed.set(query=query).render_as_string(hide_password=False))
    monkeypatch.setenv("MADRID_SETUP_TOKEN_HASH", hashlib.sha256(INVITATION.encode()).hexdigest())
    monkeypatch.setenv("MADRID_DATA_DIR", str(tmp_path / "instance1"))
    try:
        with TestClient(create_app(), headers=HEADERS) as client:
            assert client.get("/api/auth/status").json()["cloud"] is True
            assert client.post("/api/auth/setup", json={"password": PASSWORD}).status_code == 403
            assert (
                client.post(
                    "/api/auth/setup", json={"password": PASSWORD}, headers={"X-Madrid-Setup": INVITATION}
                ).status_code
                == 200
            )
            assert client.post("/api/auth/login", json={"password": PASSWORD}).status_code == 200
            yield client
    finally:
        with engine.begin() as connection:
            connection.execute(text(f'DROP SCHEMA "{schema}" CASCADE'))
        engine.dispose()


def test_postgres_keeps_original_and_batch_after_instance_replacement(cloud, tmp_path, monkeypatch):
    batch = upload(cloud, SNAPSHOT, reservation_verified=True)
    accept(cloud, batch)
    duplicate = upload(cloud, SNAPSHOT, reservation_verified=True)
    assert duplicate["id"] == batch["id"]
    with cloud.app.state.factory() as db:
        original = db.scalar(select(Original))
        assert original.content == csv_bytes(SNAPSHOT)
    assert not (cloud.app.state.directory / "originals").exists()
    assert not list(cloud.app.state.directory.glob("parse-*"))
    monkeypatch.setenv("MADRID_DATA_DIR", str(tmp_path / "replacement"))
    with TestClient(create_app(), headers=HEADERS) as replacement:
        assert replacement.get("/api/auth/status").json()["initialized"] is True
        assert replacement.post("/api/auth/login", json={"password": PASSWORD}).status_code == 200
        report = replacement.get("/api/analytics?start=2026-10-01&end=2026-10-24").json()
        assert report["metrics"]["stock"]["value"] == "10.00"
        assert replacement.get("/api/batches").json()[0]["id"] == batch["id"]


def test_postgres_workbook_sheet_selection_reconciliation_and_persistence(cloud, tmp_path, monkeypatch):
    book = Workbook()
    first = book.active
    first.title = "First"
    for row in [["barcode", "name", "quantity"], ["000001", "First item", "2"], ["TOTAL", "", "2"]]:
        first.append(row)
    second = book.create_sheet("Second")
    for row in [["quantity", "barcode", "name"], ["3", "000002", "Second item"], ["3", "", "TOTAL"]]:
        second.append(row)
    book.create_sheet("Instructions").append(["Do not import this worksheet"])
    content = io.BytesIO()
    book.save(content)
    spec = {
        "report_type": "snapshot",
        "start": "2026-10-24",
        "end": "2026-10-24",
        "sheets": ["First", "Second"],
        "complete": True,
    }
    result = cloud.post(
        "/api/imports", files={"file": ("workbook.xlsx", content.getvalue())}, data={"spec": json.dumps(spec)}
    )
    assert result.status_code == 200, result.text
    batch = result.json()
    assert not batch["issues"]
    assert batch["reconciliation"]["details"] == 2
    accept(cloud, batch)
    monkeypatch.setenv("MADRID_DATA_DIR", str(tmp_path / "replacement"))
    with TestClient(create_app(), headers=HEADERS) as replacement:
        assert replacement.post("/api/auth/login", json={"password": PASSWORD}).status_code == 200
        persisted = replacement.get(f"/api/batches/{batch['id']}").json()
        assert persisted["rows"] == batch["rows"]
        report = replacement.get("/api/analytics?start=2026-10-24&end=2026-10-24").json()
        assert report["metrics"]["stock"]["value"] == "5.00"


def test_postgres_serializes_acceptance_across_instances(cloud):
    first = upload(cloud, SNAPSHOT, reservation_verified=True, note="first interpretation")
    second = upload(cloud, SNAPSHOT, reservation_verified=True, note="second interpretation")
    with TestClient(create_app(), headers=HEADERS) as other:
        assert other.post("/api/auth/login", json={"password": PASSWORD}).status_code == 200
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [
                pool.submit(c.post, f"/api/batches/{b['id']}/action", json={"action": "accept"})
                for c, b in [(cloud, first), (other, second)]
            ]
            assert sorted(f.result().status_code for f in futures) == [200, 409]
    with cloud.app.state.factory() as db:
        assert db.scalar(select(func.count()).select_from(Batch).where(Batch.status == "accepted")) == 1


def test_postgres_original_retention_is_shared_and_audited(cloud):
    batch = upload(cloud, SNAPSHOT, reservation_verified=True)
    accept(cloud, batch)
    url = f"/api/batches/{batch['id']}/original"
    assert cloud.request("DELETE", url, json={"reason": "synthetic retention"}).status_code == 409
    assert (
        cloud.post(
            f"/api/batches/{batch['id']}/action",
            json={"action": "rollback", "reason": "synthetic correction"},
        ).status_code
        == 200
    )
    assert cloud.request("DELETE", url, json={"reason": "synthetic retention"}).status_code == 200
    with cloud.app.state.factory() as db:
        assert db.scalar(select(Original)) is None
        assert db.scalar(select(Audit).where(Audit.action == "original_deleted")) is not None
        assert db.get(Batch, batch["id"]).rows


def test_login_throttle_survives_multiple_instances_and_expiry(cloud):
    with TestClient(create_app(), headers=HEADERS) as other:
        for index in range(10):
            c = cloud if index % 2 else other
            assert (
                c.post("/api/auth/login", json={"password": "wrong-cloud-password-only"}).status_code == 401
            )
        assert other.post("/api/auth/login", json={"password": PASSWORD}).status_code == 429
        assert cloud.post("/api/auth/login", json={"password": PASSWORD}).status_code == 429
        with cloud.app.state.factory() as db:
            db.get(AuthThrottle, 1).attempts = [time.time() - 301]
            db.commit()
        assert other.post("/api/auth/login", json={"password": PASSWORD}).status_code == 200


def test_cloud_backup_restores_exact_rows_and_originals_to_local_store(cloud, tmp_path):
    batch = upload(cloud, SNAPSHOT, reservation_verified=True)
    accept(cloud, batch)
    password = "synthetic-backup-passphrase-only"
    archive = create_cloud_backup(cloud.app.state.factory.kw["bind"], password)
    target = tmp_path / "restored"
    manifest = restore_backup(archive, target, password)
    assert len(manifest) == 2
    assert next((target / "originals").iterdir()).read_bytes() == csv_bytes(SNAPSHOT)
    restored = session_factory(make_engine(target / "madrid.sqlite3"))
    with restored() as db, cloud.app.state.factory() as source:
        assert db.get(Batch, batch["id"]).rows == source.get(Batch, batch["id"]).rows
        assert db.scalar(select(Original)).content == csv_bytes(SNAPSHOT)
    with TestClient(create_app(target), headers=HEADERS) as local:
        assert local.post("/api/auth/login", json={"password": PASSWORD}).status_code == 200
        assert (
            local.get("/api/analytics?start=2026-10-01&end=2026-10-24").json()["metrics"]["stock"]["value"]
            == "10.00"
        )


def test_cloud_invitation_cannot_replace_an_existing_owner(cloud):
    response = cloud.post(
        "/api/auth/setup",
        json={"password": "replacement-password-only"},
        headers={"X-Madrid-Setup": INVITATION},
    )
    assert response.status_code == 409
    assert cloud.post("/api/auth/login", json={"password": PASSWORD}).status_code == 200


def test_cloud_database_rejects_untrusted_tls_certificate(cloud, monkeypatch, tmp_path):
    # Only local integration CI uses a private CA; ordinary public cloud DBs use the system bundle.
    if not os.environ.get("MADRID_DATABASE_CA_FILE"):
        pytest.skip("Private test CA not configured; negative trust test needs its own test server")
    monkeypatch.delenv("MADRID_DATABASE_CA_FILE")
    engine = make_engine(Path(tmp_path) / "unused", os.environ["MADRID_DATABASE_URL"])
    try:
        with pytest.raises(OperationalError, match="certificate verify failed"):
            with engine.connect():
                pass
    finally:
        engine.dispose()
