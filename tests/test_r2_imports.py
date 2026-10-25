import io
import json

from openpyxl import Workbook

from tests.test_imports import upload


def test_explicit_worksheet_modes_cannot_silently_fall_back_to_all(client):
    base = {"report_type": "snapshot", "start": "2026-10-24", "end": "2026-10-24", "complete": True}
    for invalid in [{"sheet_mode": "selected", "sheets": []}, {"sheet_mode": "all", "sheets": ["First"]}]:
        result = client.post(
            "/api/imports",
            files={"file": ("unused.xlsx", b"synthetic")},
            data={"spec": json.dumps({**base, **invalid})},
        )
        assert result.status_code == 422
    book = Workbook()
    book.active.title = "First"
    book.active.append(["barcode", "name", "quantity"])
    book.active.append(["001", "First", "2"])
    book.active.append(["TOTAL", "", "2"])
    other = book.create_sheet("Second")
    other.append(["quantity", "barcode", "name"])
    other.append(["99", "002", "Second"])
    other.append(["99", "", "TOTAL"])
    book.create_sheet("Empty")
    stream = io.BytesIO()
    book.save(stream)
    spec = {**base, "sheet_mode": "selected", "sheets": ["First"]}
    for endpoint in ["/api/imports/inspect", "/api/imports"]:
        result = client.post(
            endpoint, files={"file": ("sheets.xlsx", stream.getvalue())}, data={"spec": json.dumps(spec)}
        )
        assert result.status_code == 200, result.text
        if endpoint.endswith("inspect"):
            assert result.json()["worksheet_names"] == ["First", "Second", "Empty"]
            assert [item["name"] for item in result.json()["sheets"]] == ["First"]
        else:
            assert not result.json()["issues"]
            assert result.json()["reconciliation"]["totals"]["quantity"] == "2"


def test_intermediate_subtotal_cannot_hide_behind_correct_grand_total(client):
    batch = upload(
        client,
        [
            ["barcode", "name", "quantity"],
            ["001", "First", "2"],
            ["SUBTOTAL", "", "3"],
            ["002", "Second", "3"],
            ["SUBTOTAL", "", "2"],
            ["TOTAL", "", "5"],
        ],
    )
    assert "source_total_mismatch" in batch["issues"]
    checks = batch["reconciliation"]["checks"]
    assert checks["source_subtotal:row=3:quantity"]["difference"] == "-1"
    assert client.post(f"/api/batches/{batch['id']}/action", json={"action": "accept"}).status_code == 409


def test_formula_footer_cannot_be_trusted_as_independent_evidence(client):
    book = Workbook()
    book.active.append(["barcode", "name", "quantity"])
    book.active.append(["001", "Item", "2"])
    book.active.append(["TOTAL", "", "=SUM(C2:C2)"])
    stream = io.BytesIO()
    book.save(stream)
    result = client.post(
        "/api/imports",
        files={"file": ("formula.xlsx", stream.getvalue())},
        data={
            "spec": json.dumps(
                {
                    "report_type": "snapshot",
                    "start": "2026-10-24",
                    "end": "2026-10-24",
                    "complete": True,
                    "expected": {"quantity": "2"},
                }
            )
        },
    )
    assert result.status_code == 200
    assert "row_issues_require_corrected_reimport" in result.json()["issues"]


def test_inspection_lists_independent_sheet_headers_without_staging(client):
    book = Workbook()
    book.active.title = "Stock"
    book.active.append(["barcode", "name", "quantity"])
    book.active.append(["001", "Item", "2"])
    sheet = book.create_sheet("Other")
    sheet.append(["Description", "Balance", "Identifier"])
    sheet.append(["Other", "3", "002"])
    stream = io.BytesIO()
    book.save(stream)
    result = client.post("/api/imports/inspect", files={"file": ("stock.xlsx", stream.getvalue())})
    assert result.status_code == 200, result.text
    report = result.json()
    assert report["status"] == "NEEDS_MAPPING"
    assert report["sheets"][0]["report_type"] == "snapshot"
    assert report["sheets"][1]["columns"] == ["Description", "Balance", "Identifier"]
    assert report["sheets"][0]["preview"][0][0] == "001"
    assert client.get("/api/batches").json() == []


def test_inspection_invalid_signature_is_actionable_and_authenticated(client):
    result = client.post("/api/imports/inspect", files={"file": ("wrong.xlsx", b"not a workbook")})
    assert result.status_code == 200
    assert result.json()["status"] == "INVALID"
    assert result.json()["issues"] == ["invalid_xlsx_signature"]
    client.post("/api/auth/logout")
    assert client.post("/api/imports/inspect", files={"file": ("data.csv", b"a,b\n1,2")}).status_code == 401


def test_bad_numeric_cell_has_column_value_and_correction_context(client):
    batch = upload(
        client, [["barcode", "name", "quantity"], ["001", "Item", "1,25"]], expected={"quantity": "1.25"}
    )
    record = next(row for row in batch["rows"] if row["kind"] == "detail")
    assert record["diagnostics"][0]["column"] == "quantity"
    assert record["diagnostics"][0]["value"] == "1,25"
    assert record["raw"]["location"] == {"row": 2}


def test_csv_bad_quoting_fails_instead_of_merging_physical_records(client):
    result = client.post(
        "/api/imports/inspect",
        files={"file": ("broken.csv", b'barcode,name,quantity\n001,"unclosed,2\n002,Item,3\n')},
        data={"spec": json.dumps({"delimiter": ","})},
    )
    assert result.json()["status"] == "INVALID"
    assert result.json()["issues"] == ["invalid_csv_quoting_or_field_size"]


def test_xls_dates_become_iso_without_treating_serials_as_dates_in_other_columns(client):
    from datetime import datetime

    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("Expenses")
    for index, value in enumerate(["date", "category", "amount"]):
        sheet.write(0, index, value)
    sheet.write(1, 0, datetime(2026, 10, 24), xlwt.easyxf(num_format_str="DD/MM/YYYY"))
    sheet.write(1, 1, "Rent")
    sheet.write(1, 2, 25)
    stream = io.BytesIO()
    book.save(stream)
    result = client.post(
        "/api/imports",
        files={"file": ("expenses.xls", stream.getvalue())},
        data={
            "spec": json.dumps(
                {
                    "report_type": "expenses",
                    "start": "2026-10-24",
                    "end": "2026-10-24",
                    "complete": True,
                    "financial_verified": True,
                    "expected": {"amount": "25"},
                }
            )
        },
    )
    assert result.status_code == 200, result.text
    assert result.json()["issues"] == []
    assert result.json()["rows"][1]["data"]["date"] == "2026-10-24"


def test_printed_net_footer_supplies_independent_money_total_without_manual_reentry(client):
    batch = upload(
        client,
        [
            ["barcode", "name", "sold", "returned", "net"],
            ["001", "Item", "2", "1", "75.25"],
            ["TOTAL", "", "2", "1", "75.25"],
        ],
        kind="sales",
        financial_verified=True,
    )
    assert batch["issues"] == []
    assert batch["reconciliation"]["checks"]["net"] == {
        "expected": "75.25",
        "actual": "75.25",
        "matches": True,
        "difference": "0.00",
    }
